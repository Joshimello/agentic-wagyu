"""Lasting edge-case checks using temporary derivatives of the supplied documents."""

import os, tempfile, unittest
from pathlib import Path
from uuid import uuid4

PRIVATE = Path(os.environ.get("DATA_ROOT", "/data"))
root = PRIVATE / "tests" / f"documents-{uuid4()}"
os.environ["DATA_ROOT"] = str(root)
from documents import extract, render_page, command
from PIL import Image
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement


class Documents(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        root.mkdir(parents=True)
        cls.tmp = tempfile.TemporaryDirectory(dir=root)
        cls.directory = Path(cls.tmp.name)
        cls.pdf = PRIVATE / "documents/260629_VLM.pdf"
        cls.image = render_page(cls.pdf, 1, cls.directory)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()
        __import__("shutil").rmtree(root)

    def test_hidden_slide_notes_fonts_and_static_media_coverage(self):
        deck = Presentation(PRIVATE / "documents/260629_VLM.pptx")
        for slide_id in list(deck.slides._sldIdLst)[2:]:
            deck.part.drop_rel(slide_id.rId)
            deck.slides._sldIdLst.remove(slide_id)
        slide = deck.slides[0]
        slide._element.set("show", "0")
        slide.notes_slide.notes_text_frame.text = (
            "Separate presenter note for regression verification"
        )
        slide._element.append(OxmlElement("p:timing"))
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Missing Regression Font"
        path = self.directory / "edge.pptx"
        deck.save(path)
        result = extract(path, "pptx")
        self.assertEqual(len(result["units"]), 2)
        self.assertTrue(result["units"][0]["hidden"])
        self.assertFalse(result["units"][0]["complete"])
        self.assertTrue(any("Missing Regression Font" in w for w in result["warnings"]))
        self.assertTrue(
            any(
                t["role"] == "notes" and "Separate presenter note" in t["text"]
                for t in result["units"][0]["texts"]
            )
        )
        self.assertEqual(extract(path, "pptx")["units"], result["units"])
        self.assertTrue(extract(path, "pptx")["cached"])

    def test_scanned_pdf_ocr_has_render_and_partial_coverage(self):
        path = self.directory / "scanned.pdf"
        self.image.save(path, "PDF", resolution=150)
        result = extract(path, "pdf")
        unit = result["units"][0]
        self.assertFalse(unit["complete"])
        self.assertTrue(unit["warnings"])
        self.assertTrue(unit["texts"])
        self.assertTrue(all(t["role"] == "ocr" for t in unit["texts"]))
        self.assertTrue((root / unit["render"]["storageKey"]).is_file())

    def test_rotated_image_preserves_orientation_transform(self):
        path = self.directory / "rotated.jpg"
        rotated = self.image.rotate(90, expand=True)
        exif = rotated.getexif()
        exif[274] = 6
        rotated.save(path, exif=exif)
        result = extract(path, "image")
        transform = result["transform"]
        self.assertEqual(transform["exifOrientation"], 6)
        self.assertEqual(
            (transform["displayWidth"], transform["displayHeight"]), self.image.size
        )
        self.assertEqual(
            (transform["sourceWidth"], transform["sourceHeight"]), rotated.size
        )
        self.assertFalse(result["units"][0]["complete"])

    def test_pdf_columns_have_bounded_native_blocks_and_complete_page_mapping(self):
        # This supplied report contains tables/figures; every page remains viewable regardless of text ordering.
        result = extract(PRIVATE / "documents/ITRI_VLM_Midterm_Report.pdf", "pdf")
        self.assertGreater(len(result["units"]), 1)
        bounds = []
        for n, unit in enumerate(result["units"], 1):
            self.assertEqual(unit["number"], n)
            self.assertTrue((root / unit["render"]["storageKey"]).is_file())
            for text in unit["texts"]:
                if text["role"] == "native":
                    box = text["bounds"]
                    self.assertEqual(box["units"], "points")
                    self.assertGreater(box["canvasWidth"], 0)
                    bounds.append(box)
        self.assertGreater(len(bounds), 10)

    def test_two_columns_retain_separate_source_coordinates(self):
        # Reflow actual source paragraphs into a temporary two-column PDF, not a new corpus upload.
        source_text = command(["pdftotext", str(self.pdf), "-"]).decode()
        fragments = [
            line.strip() for line in source_text.splitlines() if len(line.strip()) > 15
        ][:2]
        self.assertEqual(len(fragments), 2)
        deck = Presentation(PRIVATE / "documents/260629_VLM.pptx")
        for slide_id in list(deck.slides._sldIdLst)[1:]:
            deck.part.drop_rel(slide_id.rId)
            deck.slides._sldIdLst.remove(slide_id)
        slide = deck.slides[0]
        for shape in list(slide.shapes):
            slide.shapes._spTree.remove(shape._element)
        for i, text in enumerate(fragments):
            shape = slide.shapes.add_textbox(
                int(deck.slide_width * (0.05 + i * 0.55)),
                int(deck.slide_height * 0.1),
                int(deck.slide_width * 0.35),
                int(deck.slide_height * 0.8),
            )
            shape.text_frame.text = text
        path = self.directory / "columns.pptx"
        deck.save(path)
        command(
            [
                "libreoffice",
                f"-env:UserInstallation={(self.directory / 'columns-profile').as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(self.directory),
                str(path),
            ]
        )
        result = extract(self.directory / "columns.pdf", "pdf")
        boxes = [
            t["bounds"] for t in result["units"][0]["texts"] if t["role"] == "native"
        ]
        self.assertTrue(any(b["x"] < b["canvasWidth"] * 0.4 for b in boxes))
        self.assertTrue(any(b["x"] > b["canvasWidth"] * 0.5 for b in boxes))
        self.assertTrue(all(b["units"] == "points" for b in boxes))

    def test_invalid_format_is_rejected(self):
        with self.assertRaises(ValueError):
            extract(self.pdf, "exe")


if __name__ == "__main__":
    unittest.main()
