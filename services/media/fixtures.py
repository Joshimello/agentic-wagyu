"""Generate reproducible synthetic media; no generated files belong in Git."""
import json, subprocess, wave
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

root = Path('/data/fixtures'); root.mkdir(parents=True, exist_ok=True)
pilot = json.loads(Path('/evals/pilot.json').read_text())
font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 44)
slides = [('LAUNCH PLAN', 'Launch: Friday', '#16324d'), ('PILOT RESULTS', 'Customers: 12', '#164d39'), ('STATUS', 'Status: green', '#164d39')]
deck = Presentation()
for n, (title, body, bg) in enumerate(slides):
    img = Image.new('RGB',(960,540),bg); draw=ImageDraw.Draw(img)
    draw.text((55,70),title,font=font,fill='white'); draw.text((55,240),body,font=font,fill='white')
    img.save(root/f'slide-{n}.png')
    slide=deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text=title; slide.placeholders[1].text=body
deck.save(root/'meeting.pptx')
samples=[]; gold=[]; rate=22050
for n,turn in enumerate(pilot['turns']):
    path=root/f'turn-{n}.wav'
    subprocess.run(['espeak-ng','-v',turn['voice'],'-s','145','-w',str(path),turn['text']],check=True)
    with wave.open(str(path),'rb') as w:
        assert w.getframerate()==rate and w.getnchannels()==1
        speech=w.readframes(w.getnframes())
    silence=b'\0\0'*rate
    samples.append(silence); start=sum(len(s) for s in samples)//2/rate
    samples.append(speech); end=sum(len(s) for s in samples)//2/rate
    gold.append({**turn,'startMs':round(start*1000),'endMs':round(end*1000)})
samples.append(b'\0\0'*rate)
with wave.open(str(root/'meeting.wav'),'wb') as w:
    w.setnchannels(1); w.setsampwidth(2); w.setframerate(rate); w.writeframes(b''.join(samples))
duration=sum(len(s) for s in samples)//2/rate
(root/'gold.json').write_text(json.dumps({'durationMs':round(duration*1000),'turns':gold},indent=2))
concat=root/'slides.txt'
concat.write_text(''.join(f"file '{root}/slide-{i}.png'\nduration {duration/3}\n" for i in range(3))+f"file '{root}/slide-2.png'\n")
subprocess.run(['ffmpeg','-y','-v','error','-f','concat','-safe','0','-i',str(concat),'-i',str(root/'meeting.wav'),'-vf','fps=10','-c:v','libx264','-pix_fmt','yuv420p','-c:a','aac','-t',str(duration),str(root/'meeting.mp4')],check=True)
subprocess.run(['ffmpeg','-y','-v','error','-stream_loop','-1','-i',str(root/'meeting.mp4'),'-t','300','-c','copy',str(root/'meeting-long.mp4')],check=True)
print(json.dumps({'fixture':str(root/'meeting.mp4'),'durationSeconds':duration,'benchmarkSeconds':300}))
