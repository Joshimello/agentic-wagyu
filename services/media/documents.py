"""Source-faithful document adapters. No source text is interpreted as instructions."""

import json, os, re, shutil, subprocess, tempfile, threading, time
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from PIL import Image, ImageOps
from pptx import Presentation
from media import ROOT, store_image
from document_media import extract_media

ADAPTER = "documents-3"
MAX_UNITS = 400
_lock = threading.Lock()  # Poppler/LibreOffice memory use is bounded to one document.


def command(args, timeout=180):
    return subprocess.run(
        args,
        check=True,
        capture_output=True,
        timeout=timeout,
        env={**os.environ, "LC_ALL": "C.UTF-8"},
    ).stdout


def versions():
    return {
        "libreoffice": command(["libreoffice", "--version"]).decode().strip(),
        "poppler": subprocess.run(
            ["pdftoppm", "-v"], capture_output=True, text=True
        ).stderr.splitlines()[0],
        "tesseract": command(["tesseract", "--version"]).decode().splitlines()[0],
        "python-pptx": __import__("pptx").__version__,
        "pillow": Image.__version__,
    }


def bbox(x, y, width, height, canvas_w, canvas_h, units):
    return {
        "x": x,
        "y": y,
        "width": width,
        "height": height,
        "canvasWidth": canvas_w,
        "canvasHeight": canvas_h,
        "units": units,
        "origin": "top_left",
    }


def ocr(img, tmp):
    target = tmp / "ocr.png"
    img.save(target)
    text = (
        command(
            [
                "tesseract",
                str(target),
                "stdout",
                "-l",
                "eng+chi_tra+chi_sim",
                "--psm",
                "3",
            ],
            120,
        )
        .decode()
        .strip()
    )
    return [{"role": "ocr", "text": text, "bounds": None}] if text else []


def render_page(pdf, number, tmp):
    prefix = tmp / "render"
    command(
        [
            "pdftoppm",
            "-f",
            str(number),
            "-l",
            str(number),
            "-singlefile",
            "-scale-to",
            "1800",
            "-jpeg",
            str(pdf),
            str(prefix),
        ]
    )
    with Image.open(prefix.with_suffix(".jpg")) as image:
        return image.convert("RGB")


def pdf_units(pdf, tmp):
    xml = command(["pdftotext", "-bbox-layout", "-enc", "UTF-8", str(pdf), "-"])
    root = ET.fromstring(xml)
    # Discard XHTML namespace; external DTDs/entities are never fetched by ElementTree.
    for node in root.iter():
        node.tag = node.tag.rsplit("}", 1)[-1]
    pages = root.findall(".//page")
    if not 1 <= len(pages) <= MAX_UNITS:
        raise ValueError(f"Document must contain 1–{MAX_UNITS} pages")
    units = []
    for n, page in enumerate(pages, 1):
        width = float(page.get("width"))
        height = float(page.get("height"))
        texts = []
        warnings = []
        for block in page.findall(".//block"):
            text = "\n".join(
                " ".join((word.text or "") for word in line.findall("word"))
                for line in block.findall("line")
            ).strip()
            if text:
                x = float(block.get("xMin"))
                y = float(block.get("yMin"))
                texts.append(
                    {
                        "role": "native",
                        "text": text,
                        "bounds": bbox(
                            x,
                            y,
                            float(block.get("xMax")) - x,
                            float(block.get("yMax")) - y,
                            width,
                            height,
                            "points",
                        ),
                    }
                )
        image = render_page(pdf, n, tmp)
        complete = True
        if not texts:
            texts = ocr(image, tmp)
            complete = False
            warnings.append(
                "No native text; OCR fallback is uncertain. The page render is authoritative, including blank pages."
            )
        units.append(
            {
                "number": n,
                "hidden": False,
                "width": image.width,
                "height": image.height,
                "texts": texts,
                "render": store_image(image),
                "complete": complete,
                "warnings": warnings,
            }
        )
    return units


def shapes_text(shapes, width, height, grouped=False):
    texts = []
    for shape in shapes:
        if shape.shape_type == 6:
            texts.extend(shapes_text(shape.shapes, width, height, True))
            continue
        values = []
        if shape.has_text_frame:
            values.append(shape.text_frame.text)
        if shape.has_table:
            values.append(
                "\n".join(
                    "\t".join(cell.text for cell in row.cells)
                    for row in shape.table.rows
                )
            )
        for text in values:
            if text.strip():
                texts.append(
                    {
                        "role": "native",
                        "text": text,
                        "bounds": None
                        if grouped
                        else bbox(
                            int(shape.left),
                            int(shape.top),
                            int(shape.width),
                            int(shape.height),
                            width,
                            height,
                            "emu",
                        ),
                    }
                )
    return texts


def pptx_units(path, tmp):
    local = tmp / "source.pptx"
    shutil.copyfile(path, local)
    presentation = Presentation(local)
    if not 1 <= len(presentation.slides) <= MAX_UNITS:
        raise ValueError(f"Document must contain 1–{MAX_UNITS} slides")
    fonts = set()
    unsupported = {}
    with ZipFile(local) as archive:
        for name in archive.namelist():
            if re.fullmatch(r"ppt/(?:slides|notesSlides)/[^/]+\.xml", name):
                content = archive.read(name)
                fonts.update(
                    f.decode()
                    for f in re.findall(rb'typeface="([^"]+)"', content)
                    if not f.startswith(b"+")
                )
        for i, slide in enumerate(presentation.slides, 1):
            tags = {node.tag.rsplit("}", 1)[-1] for node in slide._element.iter()}
            unsupported[i] = bool(
                tags & {"timing", "videoFile", "audioFile", "oleObj", "contentPart"}
            )
    available = set(
        x.strip().lower()
        for line in command(["fc-list", ":", "family"]).decode().splitlines()
        for x in line.split(",")
    )
    missing = sorted(f for f in fonts if f.lower() not in available)
    profile = tmp / "lo-profile"
    command(
        [
            "libreoffice",
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--convert-to",
            'pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"true"},"ExportNotesPages":{"type":"boolean","value":"false"}}',
            "--outdir",
            str(tmp),
            str(local),
        ],
        300,
    )
    pdf = tmp / "source.pdf"
    if not pdf.is_file():
        raise ValueError("Slide rendering failed; original preserved")
    page_info = command(["pdfinfo", str(pdf)]).decode()
    count = int(re.search(r"^Pages:\s+(\d+)", page_info, re.M).group(1))
    if count != len(presentation.slides):
        raise ValueError(
            "Rendered slide count mismatch; refusing incorrect source mappings"
        )
    units = []
    for n, slide in enumerate(presentation.slides, 1):
        texts = shapes_text(
            slide.shapes, int(presentation.slide_width), int(presentation.slide_height)
        )
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None and frame.text.strip():
                texts.append({"role": "notes", "text": frame.text, "bounds": None})
        warnings = []
        complete = True
        if unsupported[n]:
            complete = False
            warnings.append(
                "Animation, embedded media or unsupported object is represented only by a static slide render; inspect the original."
            )
        image = render_page(pdf, n, tmp)
        units.append(
            {
                "number": n,
                "hidden": slide._element.get("show", "1") in ("0", "false"),
                "width": image.width,
                "height": image.height,
                "texts": texts,
                "render": store_image(image),
                "complete": complete,
                "warnings": warnings,
            }
        )
    warnings = [
        "Static slide renders may differ from PowerPoint; charts, equations and diagrams remain in the render."
    ]
    if missing:
        warnings.append(
            "Unavailable source fonts were substituted: "
            + ", ".join(missing)
            + ". Review the renders against the original."
        )
    media = extract_media(local, presentation, units)
    return units, warnings, media


def image_units(path, tmp):
    with Image.open(path) as original:
        if getattr(original, "n_frames", 1) != 1:
            raise ValueError(
                "Multi-frame images are unsupported; upload a single image or PDF"
            )
        orientation = int(original.getexif().get(274, 1))
        source_size = list(original.size)
        image = ImageOps.exif_transpose(original).convert("RGB")
        image.thumbnail((2400, 2400))
    texts = ocr(image, tmp)
    warnings = [
        "Image text uses uncertain OCR; review the oriented image before including partial extraction."
    ]
    return [
        {
            "number": 1,
            "hidden": False,
            "width": image.width,
            "height": image.height,
            "texts": texts,
            "render": store_image(image),
            "complete": False,
            "warnings": warnings,
        }
    ], {
        "exifOrientation": orientation,
        "sourceWidth": source_size[0],
        "sourceHeight": source_size[1],
        "displayWidth": image.width,
        "displayHeight": image.height,
    }


def extract(path, format):
    if format not in ("pdf", "pptx", "image"):
        raise ValueError("Unsupported document format")
    cache = ROOT / "documents" / path.name / f"{ADAPTER}-{format}.json"
    with _lock:
        if cache.is_file():
            result = json.loads(cache.read_text())
            if all(
                (ROOT / a["storageKey"]).is_file()
                for a in [
                    *[u["render"] for u in result["units"]],
                    *[m["original"] for m in result.get("media", [])],
                    *[
                        p["asset"]
                        for m in result.get("media", [])
                        for p in m["previews"]
                    ],
                ]
            ):
                return {**result, "cached": True}
        started = time.monotonic()
        with tempfile.TemporaryDirectory() as directory:
            tmp = Path(directory)
            warnings = []
            transform = None
            media = []
            if format == "pdf":
                units = pdf_units(path, tmp)
                warnings = [
                    "Native PDF blocks preserve layout coordinates; reading order and equations/tables may require the page render. Embedded media is not interpreted."
                ]
            elif format == "pptx":
                units, warnings, media = pptx_units(path, tmp)
            else:
                units, transform = image_units(path, tmp)
        result = {
            "adapter": ADAPTER,
            "format": format,
            "units": units,
            "warnings": warnings,
            "transform": transform,
            "media": media,
            "tools": versions(),
            "elapsedMs": round((time.monotonic() - started) * 1000),
            "cached": False,
        }
        cache.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            dir=cache.parent, mode="w", delete=False
        ) as temp:
            json.dump(result, temp, ensure_ascii=False)
            temporary = Path(temp.name)
        os.replace(temporary, cache)
        return result
